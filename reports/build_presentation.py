#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/GV02-03_答辩PPT.pptx"
FIG = ROOT / "results/gv02_03/figures"

NAVY = RGBColor(18, 51, 73)
BLUE = RGBColor(34, 104, 139)
TEAL = RGBColor(31, 139, 137)
ORANGE = RGBColor(226, 132, 62)
DARK = RGBColor(39, 45, 52)
LIGHT = RGBColor(238, 244, 247)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, left, top, width, height, *, size=20, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Noto Sans SC"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_title(slide, title, number=None):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(0.16))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    if number:
        add_text(slide, str(number), 0.45, 0.42, 0.5, 0.5, size=15, bold=True, color=TEAL)
    add_text(slide, title, 0.85 if number else 0.55, 0.30, 12.0, 0.75, size=25, bold=True, color=NAVY)
    rule = slide.shapes.add_shape(1, Inches(0.55), Inches(1.08), Inches(12.2), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(205, 218, 225)
    rule.line.fill.background()


def add_bullets(slide, bullets, left=0.75, top=1.35, width=11.8, height=5.6, size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Noto Sans SC"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(14)
        paragraph.text = "• " + paragraph.text
    return box


def add_metric(slide, value, label, left, top, color=BLUE):
    shape = slide.shapes.add_shape(5, Inches(left), Inches(top), Inches(2.65), Inches(1.35))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(200, 216, 224)
    add_text(slide, value, left + 0.12, top + 0.18, 2.4, 0.55, size=27, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, label, left + 0.12, top + 0.78, 2.4, 0.36, size=12, color=DARK, align=PP_ALIGN.CENTER)


def picture(slide, path, left, top, width=None, height=None):
    kwargs = {}
    if width is not None:
        kwargs["width"] = Inches(width)
    if height is not None:
        kwargs["height"] = Inches(height)
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.25), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    add_text(slide, "GvpA 候选序列的\n多目标评价与筛选", 0.9, 1.15, 11.4, 1.8, size=35, bold=True, color=WHITE)
    add_text(slide, "机器学习综合实践 · GV02-03", 0.95, 3.25, 8.5, 0.5, size=19, color=RGBColor(184, 215, 227))
    add_text(slide, "约束 · 保守性 · 新颖性 · 多样性", 0.95, 4.25, 10.5, 0.6, size=22, bold=True, color=RGBColor(117, 218, 199))
    add_text(slide, "答辩时请补充真实成员姓名与分工", 0.95, 6.55, 8.5, 0.35, size=11, color=RGBColor(180, 190, 198))

    slide = prs.slides.add_slide(blank); add_title(slide, "任务定位：不是再训练一个生成模型", 1)
    add_bullets(slide, [
        "输入：T05 已生成的 200 条蛋白质候选；输出：逐条四维证据和固定预算筛选名单",
        "RQ1：四个目标如何冲突？  RQ2：不同聚合策略何时更合适？  RQ3：参数微调是否改变名单？",
        "没有候选功能真值：不报告准确率或虚构湿实验成功率，只评价代理指标、重叠与稳定性",
    ], top=1.45, size=21)
    add_metric(slide, "4", "评价维度", 0.9, 5.55, TEAL); add_metric(slide, "3", "聚合策略", 3.8, 5.55, BLUE)
    add_metric(slide, "4+2", "必做 + 选做实验", 6.7, 5.55, ORANGE); add_metric(slide, "42", "固定随机种子", 9.6, 5.55, TEAL)

    slide = prs.slides.add_slide(blank); add_title(slide, "数据与参考集：候选和天然参考严格分离", 2)
    add_metric(slide, "200", "T05 生成候选", 0.8, 1.5); add_metric(slide, "1224", "清洗天然参考", 3.75, 1.5, TEAL)
    add_metric(slide, "418", "CD-HIT 90% 代表", 6.7, 1.5, ORANGE); add_metric(slide, "23", "统计保守位点", 9.65, 1.5, TEAL)
    add_bullets(slide, [
        "参考要求显式 GvpA 标签；排除 GvpJ、partial/fragment、非标准残基与异常长度",
        "混合多家族 real_gvp.fasta 不作为纯 GvpA 参考；候选不参与保守位点定义",
        "所有来源、清洗原因、稳定 ID 与序列哈希可追溯",
    ], top=3.35, size=18)

    slide = prs.slides.add_slide(blank); add_title(slide, "可复现评价流水线", 3)
    stages = [("清洗/去冗余", "CD-HIT"), ("参考 MSA", "MAFFT"), ("域约束", "PF00741 + HMMER"), ("相似性", "BLAST+"), ("聚合/统计", "Python")]
    for index, (name, tool) in enumerate(stages):
        left = 0.55 + index * 2.55
        shape = slide.shapes.add_shape(5, Inches(left), Inches(2.0), Inches(2.2), Inches(1.5))
        shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = TEAL
        add_text(slide, name, left + .1, 2.25, 2.0, .4, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(slide, tool, left + .1, 2.85, 2.0, .3, size=12, color=BLUE, align=PP_ALIGN.CENTER)
        if index < len(stages)-1:
            add_text(slide, "→", left+2.2, 2.43, .35, .35, size=18, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
    add_bullets(slide, [
        "PF00741.24 的 GA domain cutoff=25 bits；参考覆盖分布校准模型覆盖门槛=0.95",
        "配置、输入/输出 SHA-256、工具和模型版本写入 manifest；错误停止，不伪造缺失分数",
    ], top=4.35, size=18)

    slide = prs.slides.add_slide(blank); add_title(slide, "四个指标：同为 [0,1]，含义不同", 4)
    add_bullets(slide, [
        "约束：HMM domain score 的参考经验百分位 × 模型覆盖因子",
        "保守性：候选在 23 个参考统计保守列上与共识一致的比例，缺口计 0",
        "新颖性：1 − 与最近天然参考的覆盖校正 identity",
        "多样性：候选级平均距离用于排序；子集平均两两距离用于评价最终集合",
    ], top=1.5, size=21)
    add_text(slide, "关键修正：PF01132 是 EFP 相关域；本实验使用 PF00741，但其命中仍不是 GvpA 功能证明。", 0.85, 5.9, 11.7, .7, size=17, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "RQ1：质量代理与新颖性存在明显冲突", 5)
    picture(slide, FIG / "correlation_heatmaps.png", 0.6, 1.25, width=8.0)
    add_text(slide, "Pearson", 9.0, 1.45, 3.4, .35, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_metric(slide, "+0.750", "约束–保守性", 9.4, 2.0, TEAL)
    add_metric(slide, "−0.796", "约束–新颖性", 9.4, 3.55, ORANGE)
    add_metric(slide, "−0.913", "保守–新颖性", 9.4, 5.1, ORANGE)

    slide = prs.slides.add_slide(blank); add_title(slide, "输入池审计：多数候选缺乏目标域证据", 6)
    picture(slide, FIG / "metric_distributions.png", 0.45, 1.15, width=8.2)
    add_metric(slide, "37/200", "PF00741 可报告命中", 9.25, 1.55, BLUE)
    add_metric(slide, "27/200", "通过 GA + 覆盖", 9.25, 3.25, ORANGE)
    add_text(slide, "这说明 T05 是混合 Gvp 生成池，\n不是已经确认的纯 GvpA 候选池。", 9.0, 5.15, 3.2, 1.0, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "RQ2：Top-20 三种策略各有侧重", 7)
    picture(slide, FIG / "strategy_quality_k20.png", 0.45, 1.18, width=8.3)
    add_bullets(slide, [
        "等权：域通过 100%，质量优先",
        "Pareto：多样性 0.8753，保留取舍",
        "轮转：多样性 0.8857，覆盖单项优势",
        "没有任何策略四维同时最优",
    ], left=9.0, top=1.65, width=3.6, height=4.9, size=17)

    slide = prs.slides.add_slide(blank); add_title(slide, "名单差异不是形式差异", 8)
    picture(slide, FIG / "strategy_overlap_k20.png", 0.7, 1.25, width=6.5)
    add_text(slide, "Top-20 两两交集", 8.1, 1.5, 3.9, .45, size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_metric(slide, "9", "等权 ∩ Pareto", 8.7, 2.2, BLUE)
    add_metric(slide, "10", "等权 ∩ 轮转", 8.7, 3.85, TEAL)
    add_metric(slide, "9", "Pareto ∩ 轮转", 8.7, 5.5, ORANGE)

    slide = prs.slides.add_slide(blank); add_title(slide, "消融：约束与保守性最不可替代", 9)
    data = [("去约束", "0.143", "约束均值→0.135"), ("去保守", "0.212", "保守均值→0.413"), ("去新颖", "0.818", "新颖均值→0.253"), ("去独特", "1.000", "当前名单不变")]
    for index, (name, jac, note) in enumerate(data):
        left = 0.65 + index*3.15
        add_text(slide, name, left, 1.55, 2.7, .4, size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_metric(slide, jac, "Top-20 Jaccard", left+.02, 2.1, ORANGE if index<2 else TEAL)
        add_text(slide, note, left, 3.8, 2.7, .6, size=14, color=DARK, align=PP_ALIGN.CENTER)
    add_text(slide, "消融反映当前等权选择机制；它不能把统计代理提升为因果证据。", 0.9, 5.55, 11.5, .7, size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "RQ3：局部权重稳定，覆盖阈值仍敏感", 10)
    picture(slide, FIG / "weight_robustness.png", 0.45, 1.35, width=7.6)
    add_metric(slide, "0.906", "Jaccard 均值", 8.7, 1.45, TEAL)
    add_metric(slide, "0.667", "Jaccard 最小", 8.7, 3.05, ORANGE)
    add_metric(slide, "0.996", "排名 Spearman 均值", 8.7, 4.65, BLUE)
    add_text(slide, "覆盖门槛 0.85 / 0.95 / 1.00 → 34 / 27 / 21 条合格", 0.8, 6.35, 11.7, .45, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank); add_title(slide, "选做：1000 次随机筛选基线", 11)
    add_metric(slide, "0.782", "随机集合多样性均值", 0.8, 1.6, BLUE)
    add_metric(slide, "0.612", "等权集合多样性", 3.75, 1.6, ORANGE)
    add_metric(slide, "0.875", "Pareto 集合多样性", 6.7, 1.6, TEAL)
    add_metric(slide, "0.886", "轮转集合多样性", 9.65, 1.6, TEAL)
    add_bullets(slide, [
        "三策略显著提高约束、保守性、独特性和域通过率",
        "Pareto/轮转的集合多样性超过随机 95th percentile 0.867",
        "三策略新颖性均低于随机：这是质量–新颖性冲突的真实代价",
    ], top=3.7, size=19)

    slide = prs.slides.add_slide(blank); add_title(slide, "结论与下游选择建议", 12)
    add_bullets(slide, [
        "质量优先：域/长度硬门槛 + 可解释加权；本次等权 Top-20 均通过当前域门槛",
        "探索优先：Pareto 或轮转，保留原始四维证据，不能把离群优势当成功能保证",
        "发布名单时同时冻结配置、资源版本和稳定性结果，不只交一个综合分",
        "后续必须由结构预测、GvpA/GvpJ 区分和表达/组装实验验证",
    ], top=1.45, size=20)
    add_text(slide, "16 项自动化测试全部通过 · 输入/输出 SHA-256 与工具版本均可审计", 0.8, 6.35, 11.7, .45, size=16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "谢谢", 0.8, 1.7, 11.7, .8, size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, "问题与讨论", 0.8, 2.8, 11.7, .6, size=25, color=RGBColor(117, 218, 199), align=PP_ALIGN.CENTER)
    add_text(slide, "答辩前请补充真实成员、分工、会议与过程证据；不要补写未发生的活动。", 1.2, 5.7, 10.9, .6, size=14, color=RGBColor(185, 202, 211), align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
